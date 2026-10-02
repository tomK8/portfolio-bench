#!/usr/bin/env python3
"""
Build the demo dataset: synthetic broker exports plus a synthetic database.

Two phases, run together by default because one is useless without the other:

  FILE PHASE     -> demo/input   Holdings and orders. Holdings have no table at all --
                                 PortfolioGatherer re-parses the broker exports on every
                                 sync -- so a demo portfolio only exists if these files
                                 exist. Orders need both halves: the table carries the
                                 closed legs of a completed ladder round trip, long after
                                 they left the broker's export, while the live Open orders
                                 tab is built purely from parsed files.

  DATABASE PHASE -> demo/db      Public market data copied verbatim; the ledger, orders and
                                 ladder/switch groups rebuilt with synthetic figures. The
                                 ledger is generated here rather than in files because the
                                 files cannot reproduce it: TRANSFER and Equate rows are
                                 made by endpoints with no CSV path, and symbols re-keyed
                                 through the Mappings tab do not survive a fresh parse.

Figures are produced by a generator profile that lives outside this repository (see
PROFILE_ENV). Prices, FX rates and percentages are never passed through it, so the demo
stays internally consistent against live market data.

Instruments are real. Sizes are not, and bear no documented relationship to anything.

Both phases rebuild from scratch -- demo/input and demo/db are wiped and regenerated. There
is no incremental path on purpose: everything in them is either derived from public market
data or written by the demo app, so a partial refresh could leave the two disagreeing.

Usage:
    python3 scripts/make-demo-data.py                 # both phases (the normal case)
    python3 scripts/make-demo-data.py --files-only    # after a fresh broker export
    python3 scripts/make-demo-data.py --db-only       # after upstream data changed
"""

import argparse
import csv
import importlib.util
import io
import os
import re
import shutil
import sqlite3
import sys
import uuid
from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required:  pip3 install openpyxl")

# The generator profile is deliberately not part of this repository and has no default
# implementation here. Without it nothing can be built.
PROFILE_ENV = "PB_DEMO_PROFILE"
PROFILE_DEFAULT = Path.home() / ".portfolio-bench-demo" / "profile.py"

# Namespace for the II filenames, which the broker delivers UUID-named and the parsers
# disambiguate by a header sniff.
NS = uuid.UUID("6f1d5a7e-0000-4000-8000-000000000001")
UUID_GLOB = "????????-????-????-????-????????????.csv"
PLACEHOLDERS = {"", "-", "n/a", "N/A", "na"}

# Currency prefix, sign, digits with optional thousands separators and decimals.
MONEY = re.compile(r'^(?P<pre>[^\d\-+]*)(?P<sign>[-+]?)(?P<num>[\d,]*\.?\d+)(?P<post>.*)$')


def load_profile():
    """Load the out-of-tree generator profile, or explain how to supply one."""
    path = Path(os.environ.get(PROFILE_ENV, PROFILE_DEFAULT)).expanduser()
    if not path.exists():
        sys.exit(
            f"No generator profile at {path}.\n"
            f"Demo data cannot be built without one. Point {PROFILE_ENV} at a Python file\n"
            f"exposing randomize(value) -> number and describe() -> str.")
    spec = importlib.util.spec_from_file_location("pb_demo_profile", path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    for fn in ("randomize", "describe"):
        if not hasattr(mod, fn):
            sys.exit(f"Generator profile at {path} does not define {fn}().")
    return mod


PROFILE = None


# =============================================================================================
#                                           SHARED                                           
# =============================================================================================

def stable_uuid(role):
    return str(uuid.uuid5(NS, role))


def _synth_number(numtext):
    """
    Produce the demo figure for a number written as text, keeping how it was written.

    Thousands separators and decimal places are preserved because both phases feed parsers
    that read these strings back: the broker exports quote "3,711.60" and the ledger
    descriptions carry a bare "7,605".
    """
    had_commas = "," in numtext
    decimals = len(numtext.split(".")[1]) if "." in numtext else 0
    value = PROFILE.randomize(float(numtext.replace(",", "")))
    return f"{value:,.{decimals}f}" if had_commas else f"{value:.{decimals}f}"


def synth_text(text):
    """Regenerate a money/quantity cell, preserving currency prefix and thousands style."""
    if text is None:
        return text
    s = str(text).strip()
    if s in PLACEHOLDERS:
        return text
    m = MONEY.match(s)
    if not m:
        return text
    try:
        body = _synth_number(m.group("num"))
    except ValueError:
        return text
    return f"{m.group('pre')}{m.group('sign')}{body}{m.group('post')}"


def synth_cell(value):
    """Regenerate a spreadsheet cell: numbers numerically, strings textually."""
    if value is None or isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return PROFILE.randomize(value)
    return synth_text(value)


# =============================================================================================
#                                         FILE PHASE                                         
# =============================================================================================

# ---------------------------------------------------------------- file plumbing

def bom_prefix(path):
    """The file's exact leading BOM bytes - II exports accumulate several, so preserve them."""
    data = path.read_bytes()
    n = 0
    while data[n:n + 3] == b"\xef\xbb\xbf":
        n += 3
    return data[:n].decode("utf-8")


def read_csv_rows(path):
    text = path.read_text(encoding="utf-8-sig", errors="replace").lstrip("﻿")
    return list(csv.reader(text.splitlines()))


def write_csv(path, rows, bom=""):
    out = io.StringIO()
    csv.writer(out, lineterminator="\r\n").writerows(rows)
    path.write_text(bom + out.getvalue(), encoding="utf-8", newline="")


def newest(paths):
    return max(paths, key=lambda p: p.stat().st_mtime) if paths else None


def find(directory, pattern, header_test=None):
    """Files matching a glob, optionally narrowed by a first-line sniff."""
    if not directory.is_dir():
        return []
    out = []
    for p in sorted(directory.glob(pattern)):
        if not p.is_file():
            continue
        if header_test is not None:
            try:
                head = p.read_text(encoding="utf-8-sig", errors="replace").split("\n", 1)[0]
            except Exception:
                continue
            if not header_test(head):
                continue
        out.append(p)
    return out


# II delivers holdings, cash and orders all as UUID-named CSVs; only the header tells them
# apart. These mirror the sniffs in IISippParser / IICashStatementParser / IIOrdersParser.
def is_ii_cash(head):
    return "Debit" in head and "Credit" in head and "Running Balance" in head


def is_ii_orders(head):
    return "Submitted" in head and "Settlement Date" in head and "Status" in head


def is_ii_holdings(head):
    return "Book Cost" in head and "Market Value" in head and not is_ii_cash(head)


def col_index(header, *names):
    norm = [str(h or "").strip().lower() for h in header]
    for n in names:
        if n.strip().lower() in norm:
            return norm.index(n.strip().lower())
    return None


def synth_columns(rows, header, synth_names):
    idxs = [i for i in (col_index(header, n) for n in synth_names) if i is not None]
    out = []
    for r in rows:
        r = list(r)
        for i in idxs:
            if i < len(r):
                r[i] = synth_text(r[i])
        out.append(r)
    return out


def find_header_row(ws, *required):
    """First row carrying all the named header cells (these sheets have a metadata preamble)."""
    for r in range(1, ws.max_row + 1):
        vals = [str(c.value or "").strip().lower() for c in ws[r]]
        if all(any(v == req.lower() or v.startswith(req.lower()) for v in vals)
               for req in required):
            return r
    return None


def synth_sheet(src, dst, log, label, required, synth_names):
    """Regenerate named columns of an xlsx that has a preamble above its header row."""
    wb = openpyxl.load_workbook(src)
    ws = wb[wb.sheetnames[0]]
    header_row = find_header_row(ws, *required)
    if header_row is None:
        log(f"  {label:17s}: SKIPPED - no header row in {src.name}")
        return
    header = [str(c.value or "").strip().lower() for c in ws[header_row]]
    wanted = {n.lower() for n in synth_names}
    targets = [i for i, h in enumerate(header) if h in wanted]
    n = 0
    for r in range(header_row + 1, ws.max_row + 1):
        if all(c.value is None for c in ws[r]):
            continue
        for i in targets:
            cell = ws.cell(row=r, column=i + 1)
            cell.value = synth_cell(cell.value)
        n += 1
    wb.save(dst)
    log(f"  {label:17s}: {n:4d} rows  {src.name} -> {dst.name}")


def synth_csv(src, dst, log, label, synth_names, bom=""):
    rows = read_csv_rows(src)
    if len(rows) < 2:
        log(f"  {label:17s}: SKIPPED - {src.name} has no data rows")
        return
    header, body = rows[0], rows[1:]
    write_csv(dst, [header] + synth_columns(body, header, synth_names), bom=bom)
    log(f"  {label:17s}: {len(body):4d} rows  {src.name} -> {dst.name}")


AJBELL_HOLDINGS = ["Quantity", "Value (£)", "Cost (£)", "Change (£)"]
II_HOLDINGS = ["Qty", "Day Gain/Loss", "Market Value £", "Market Value",
               "Book Cost", "Gain/Loss"]
ROTH_HOLDINGS = ["quantity", "market value", "gain/loss $"]
ORDER_SIZES = ["Quantity", "Qty", "Amount", "Value", "Consideration"]
ROTH_ORDER_SIZES = ["quantity", "qty", "amount", "value"]


# =============================================================================================
#                                       DATABASE PHASE                                       
# =============================================================================================

# Market data: public, shared with production.
PUBLIC_TABLES = [
    "price_history", "price_intraday", "fx_rates",
    "dividend_history", "dividend_scan",
    "earnings_dates", "earnings_calendar", "earnings_calendar_fetch",
    "fundamentals_snapshot", "ticker_mappings", "instrument_aliases",
    # Tickers being watched and their alert thresholds. No position size, no balance -
    # carried so the Watchlist tab works, consistent with the demo keeping real instruments.
    "watchlist",
]

# Position data: rebuilt with synthetic figures. Columns listed are regenerated;
# everything else (prices, FX rates, currencies, dates, statuses, references) is verbatim.
SYNTH_TABLES = {
    "cash_transactions": ["quantity", "amount", "amount_gbp",
                          "cash_balance", "cash_balance_gbp"],
    "dividend_events": ["dividend_amount", "dividend_gbp"],
    "order_history": ["quantity", "amount"],
    "order_group": [],
    "order_group_member": [],
    "order_group_intent": ["quantity"],
}

# Never copied: nothing reads portfolio_snapshots, and trade_notes is the journal.
OMITTED_TABLES = ["portfolio_snapshots", "trade_notes"]

# KeyValueStore is one <key>.txt file per key beside portfolio.db, not a table.
KV_VERBATIM = [
    "gilt_sedols", "earnings_checked", "unresolved_tickers", "risk_free_rate",
    "watchlist_symbols", "fundamentals_extra_symbols", "whatif_substitutes",
    "momentum_symbols", "momentum_signals_symbols", "momentum_universe",
    "momentum_universe_history", "swing_symbols", "swing_cluster",
    "trend_symbols", "range_symbols", "orders_use_drift", "orders_use_event_var",
]
# One per II SIPP cash box on the Holdings form (IiCashField) - a missing one would show the
# demo an empty box, which reads as "keep the stored balance", not as zero.
KV_SYNTH = ["roth_balance_brought_forward", "ii_sipp_cash_last", "ii_sipp_cash_last_usd",
            "ii_sipp_cash_last_eur"]
# Derived from the ledger on import/sync - copying them would leak a real balance. The copy above
# is an allowlist, so this list documents what is left behind rather than enforcing it.
KV_NEVER = ["cash_balance_II_GBP_gbp", "cash_balance_II_GBP_native",
            "cash_balance_II_USD_gbp", "cash_balance_II_USD_native",
            "cash_balance_II_EUR_gbp", "cash_balance_II_EUR_native",
            "held_symbols", "order_symbols", "watchlist_alerts_fired",
            "allocation_targets"]


# Quantities embedded in a ledger description, so a row reads consistently with the
# columns beside it. Cosmetic only - the importers read columns, not prose - but a demo where the
# description disagrees with the quantity beside it is a demo with a visible bug.
_DESC_PATTERNS = [
    re.compile(r'(?i)^(dividend\s+)([\d,]+(?:\.\d+)?)(\s+)'),
    re.compile(r'(?i)^((?:purchase|sale|redemption)\s+)([\d,]+(?:\.\d+)?)(\s+)'),
    re.compile(r'(?i)^(div\s+)(\d[\d,]*(?:\.\d+)?)(\s+)'),
    re.compile(r'(?i)^((?:buy|sell)\s+)(-?[\d,]*\.?\d+)(\s+share)'),
    re.compile(r'^(\d[\d,]*(?:\.\d+)?)(\s+)()'),
]


def synth_description(desc):
    if not isinstance(desc, str) or not desc:
        return desc
    for pat in _DESC_PATTERNS:
        m = pat.match(desc)
        if not m:
            continue
        groups = list(m.groups())
        for i, g in enumerate(groups):
            try:
                float(str(g).replace(",", ""))
            except (ValueError, TypeError):
                continue
            groups[i] = _synth_number(g)
            return desc[:m.start()] + "".join(groups) + desc[m.end():]
    return desc


def plain(value):
    """Mirror Java BigDecimal.stripTrailingZeros().toPlainString()."""
    d = Decimal(str(value)).normalize()
    if d == d.to_integral_value():
        d = d.quantize(Decimal(1))
    return format(d, "f")


def demo_order_key(prod_key):
    """
    Translate an upstream order_key into the one the demo produces.

    OpenOrder.identity() prefers broker|reference, which carries no size and is therefore
    unchanged. Only the reference-less 6-part composite embeds the quantity - Roth IRA
    orders - so only that form needs regenerating. Getting this wrong would orphan
    every ladder rung, because order_group_member joins on exactly this string.
    """
    if not prod_key:
        return prod_key
    parts = prod_key.split("|")
    if len(parts) != 6 or not parts[3]:
        return prod_key
    try:
        parts[3] = plain(PROFILE.randomize(float(parts[3])))
    except ValueError:
        return prod_key
    return "|".join(parts)


# =============================================================================================
#                                           PHASES                                           
# =============================================================================================

def build_files(downloads, out, log):
    """Regenerate the holdings and order exports into demo/input."""
    if out.exists():
        shutil.rmtree(out)
    out.mkdir(parents=True)

    log(f"FILE PHASE  ->  {out}")
    log(f"  source: {downloads}  (read-only)\n")
    missing = []

    def do(label, candidates, fn):
        f = newest(candidates)
        if f:
            fn(f)
        else:
            missing.append(label)

    do("AJ Bell holdings (portfolio*.csv)", find(downloads, "portfolio*.csv"),
       lambda f: synth_csv(f, out / "portfolio-DEMO-SIPP.csv", log,
                           "AJ Bell holdings", AJBELL_HOLDINGS, bom="﻿"))
    do("Roth holdings (Holdings*.xlsx)", find(downloads, "Holdings*.xlsx"),
       lambda f: synth_sheet(f, out / "Holdings.xlsx", log, "Roth holdings",
                             ("security id", "quantity"), ROTH_HOLDINGS))
    do("II holdings (UUID csv)", find(downloads, UUID_GLOB, is_ii_holdings),
       lambda f: synth_csv(f, out / f"{stable_uuid('ii-holdings')}.csv", log,
                           "II holdings", II_HOLDINGS, bom=bom_prefix(f)))

    do("AJ Bell orders (activeorders*.csv)", find(downloads, "activeorders*.csv"),
       lambda f: synth_csv(f, out / "activeorders_demo.csv", log,
                           "AJ Bell orders", ORDER_SIZES, bom="﻿"))
    do("Roth orders (Order+Status*.xlsx)", find(downloads, "*rder*tatus*.xlsx"),
       lambda f: synth_sheet(f, out / "Order+Status_demo.xlsx", log, "Roth orders",
                             ("date", "quantity"), ROTH_ORDER_SIZES))
    do("II orders (UUID csv)", find(downloads, UUID_GLOB, is_ii_orders),
       lambda f: synth_csv(f, out / f"{stable_uuid('ii-orders')}.csv", log,
                           "II orders", ORDER_SIZES, bom=bom_prefix(f)))

    if missing:
        log("\n  not generated (source export absent from Downloads):")
        for m in missing:
            log(f"    - {m}")
    log(f"\n  {len(list(out.iterdir()))} files written")


def build_database(source, out, intraday_days, log):
    """Build demo/db: market data verbatim, position data regenerated."""
    if not source.exists():
        sys.exit(f"No production database at {source}")

    # Wipe and rebuild rather than refresh in place. demo/db holds more than portfolio.db:
    # the KeyValueStore .txt files, and archived broker exports the demo app moved in on an
    # import. Several of those the app writes itself (held_symbols, order_symbols, and after
    # an Import cash, cash_balance_II_*) and this script never overwrites, because copying
    # them from production would carry a real balance across. Left in place they outlive the
    # data they describe -- and latestCashBalances PREFERS the asserted cash_balance_II_*
    # over any stored ledger row, so a leftover would quietly override a freshly built
    # ledger, from a run whose figures no longer apply.
    if out.exists():
        shutil.rmtree(out)
    out.mkdir(parents=True)
    target = out / "portfolio.db"
    uri = f"file:{source}?immutable=1"
    src = sqlite3.connect(uri, uri=True)
    dst = sqlite3.connect(target)
    cutoff = (date.today() - timedelta(days=intraday_days)).isoformat()

    log(f"DATABASE PHASE  ->  {target}")
    log(f"  source: {source}  (read-only, immutable)\n")
    try:
        ddl = {r[0]: r[1] for r in src.execute(
            "SELECT name, sql FROM sqlite_master WHERE type='table' AND sql IS NOT NULL")}
        index_ddl = [(r[0], r[1], r[2]) for r in src.execute(
            "SELECT name, tbl_name, sql FROM sqlite_master "
            "WHERE type='index' AND sql IS NOT NULL")]

        for table in PUBLIC_TABLES:
            if table not in ddl:
                log(f"  {table:24s} absent in source - skipped")
                continue
            dst.execute(ddl[table])
            cols = [r[1] for r in src.execute(f"PRAGMA table_info({table})")]
            collist = ",".join(f'"{c}"' for c in cols)
            where = ""
            if table == "price_intraday":
                where = f" WHERE ts >= '{cutoff}'"
            rows = list(src.execute(f"SELECT {collist} FROM {table}{where}"))
            if rows:
                marks = ",".join("?" * len(cols))
                dst.executemany(f"INSERT INTO {table} ({collist}) VALUES ({marks})", rows)
            dst.commit()
            note = f"  (pruned to {intraday_days}d)" if where else ""
            log(f"  {table:24s} {len(rows):>7,} rows{note}")

        log("")
        for table, synth_cols in SYNTH_TABLES.items():
            if table not in ddl:
                log(f"  {table:24s} absent in source - skipped")
                continue
            dst.execute(ddl[table])
            cols = [r[1] for r in src.execute(f"PRAGMA table_info({table})")]
            collist = ",".join(f'"{c}"' for c in cols)
            rows = [list(r) for r in src.execute(f"SELECT {collist} FROM {table}")]

            idx = {c: i for i, c in enumerate(cols)}
            for r in rows:
                for c in synth_cols:
                    i = idx.get(c)
                    if i is not None and isinstance(r[i], (int, float)):
                        r[i] = PROFILE.randomize(r[i])
                if "description" in idx:
                    r[idx["description"]] = synth_description(r[idx["description"]])
                # order_key embeds the quantity when the broker gave no reference, so the
                # demo produces a different string. Rebuild it on both sides of the
                # join or every ladder rung is orphaned.
                if "order_key" in idx:
                    r[idx["order_key"]] = demo_order_key(r[idx["order_key"]])
            if rows:
                marks = ",".join("?" * len(cols))
                dst.executemany(f"INSERT INTO {table} ({collist}) VALUES ({marks})", rows)
            dst.commit()
            note = f"  ({', '.join(synth_cols)})" if synth_cols else ""
            log(f"  {table:24s} {len(rows):>7,} rows{note}")

        for name, tbl, sql in index_ddl:
            if tbl in PUBLIC_TABLES or tbl in SYNTH_TABLES:
                try:
                    dst.execute(sql)
                except sqlite3.OperationalError:
                    pass
        dst.commit()
        dst.execute("VACUUM")
    finally:
        src.close()
        dst.close()

    log(f"\n  left empty on purpose: {', '.join(OMITTED_TABLES)}")

    # ---- KeyValueStore files
    src_dir = source.parent
    copied, synth_n = 0, 0
    for key in KV_VERBATIM:
        f = src_dir / f"{key}.txt"
        if f.exists():
            shutil.copy2(f, out / f.name)
            copied += 1
    for key in KV_SYNTH:
        f = src_dir / f"{key}.txt"
        if not f.exists():
            continue
        try:
            value = float(f.read_text().strip())
        except ValueError:
            continue
        (out / f.name).write_text(f"{PROFILE.randomize(value):.2f}")
        synth_n += 1
    log(f"  key-value files: {copied} copied, {synth_n} regenerated, "
          f"{len(KV_NEVER)} deliberately omitted")

    size_mb = target.stat().st_size / (1024 * 1024)
    log(f"\nDone. {target}  ({size_mb:.0f} MB)")


def demo_is_running(root):
    """PID of a live demo instance, or None. Wiping demo/db under one would corrupt it."""
    pid_file = root / "portfolio-bench-demo.pid"
    if not pid_file.exists():
        return None
    try:
        pid = int(pid_file.read_text().strip())
        os.kill(pid, 0)
        return pid
    except (ValueError, ProcessLookupError, PermissionError, OSError):
        return None


def main():
    home = Path.home()
    root = Path(__file__).resolve().parent.parent
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--downloads", type=Path, default=home / "Downloads",
                    help="where the real broker exports live (read-only)")
    ap.add_argument("--source", type=Path,
                    default=home / "Documents" / "Investing" / "portfolio.db",
                    help="production database (read-only)")
    ap.add_argument("--out", type=Path, default=root / "demo",
                    help="demo directory; input/ and db/ are created under it")
    ap.add_argument("--intraday-days", type=int, default=14,
                    help="keep this many days of intraday bars (default: 14)")
    phase = ap.add_mutually_exclusive_group()
    phase.add_argument("--files-only", action="store_true",
                       help="rebuild demo/input only, leaving the database alone")
    phase.add_argument("--db-only", action="store_true",
                       help="rebuild demo/db only, leaving the exports alone")
    args = ap.parse_args()

    do_files = not args.db_only
    do_db = not args.files_only

    if do_db:
        pid = demo_is_running(root)
        if pid is not None:
            sys.exit(f"The demo is running (PID {pid}) and this rebuilds its database.\n"
                     f"Stop it first:  scripts/demo-stop.sh\n"
                     f"Or do both:     scripts/demo-reset.sh")

    global PROFILE
    PROFILE = load_profile()

    def log(msg):
        print(msg)

    log(f"generator: {PROFILE.describe()}\n")

    if do_files:
        build_files(args.downloads, args.out / "input", log)
    if do_files and do_db:
        log("")
    if do_db:
        build_database(args.source, args.out / "db", args.intraday_days, log)

    if do_files and do_db:
        log("\nNext: scripts/demo.sh")


if __name__ == "__main__":
    main()
